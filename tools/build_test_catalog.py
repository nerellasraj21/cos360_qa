import csv
import glob
import json
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPECS = os.path.join(ROOT, "test_specs")
REPORTS = os.path.join(ROOT, "reports")
OUT_XLSX = os.path.join(ROOT, "test_cases", "COS360_Test_Cases.xlsx")
OUT_CSV = os.path.join(ROOT, "test_cases", "COS360_Test_Cases.csv")

PHASES = {"U": "Unit", "A": "API", "E": "UI"}
ID_RE = re.compile(r"TC-[A-Z]+-\d\d-[UAE]\d\d")
FEATURE_RE = re.compile(r"^##\s+(F\d\d)\s+(.*)$")
ROW_RE = re.compile(r"^\|\s*(TC-[A-Z]+-\d\d-[UAE]\d\d)\s*\|")
JOURNEY_RE = re.compile(r"^##\s+(J\d\d)\s+(.*)$")
STEP_RE = re.compile(r"^(\d+)\.\s+(.*)$")
MODULE_TITLES = {
    "auth": "Authentication",
    "tenants-and-admin": "Tenants and Admin",
    "masters": "Masters",
    "timetable-calendar": "Timetable and Calendar",
    "students": "Students",
    "certificates": "Certificates",
    "staff": "Staff",
    "communication": "Communication",
    "fee": "Fee",
    "exam": "Exam",
    "expense": "Expense",
    "transport": "Transport",
    "reports-dashboards": "Reports and Dashboards",
}
MANUAL_FIELDS = ["result", "actual", "defect", "tester", "test_date", "build"]
RESULT_VALUES = '"Pass,Fail,Blocked,Not run"'

CASE_COLUMNS = [
    ("id", "Test Case ID", 17),
    ("module", "Module", 16),
    ("feature", "Feature", 26),
    ("phase", "Type", 7),
    ("priority", "Priority", 8),
    ("platform", "Platform", 9),
    ("role", "Role", 10),
    ("preconditions", "Preconditions", 34),
    ("steps", "Scenario / Steps", 60),
    ("expected", "Expected Result", 50),
    ("doc_status", "Doc Status", 18),
    ("automated", "Automated", 10),
    ("last_run", "Last Run", 10),
    ("result", "Result", 10),
    ("actual", "Actual Result", 36),
    ("defect", "Defect ID", 12),
    ("tester", "Tester", 12),
    ("test_date", "Test Date", 12),
    ("build", "Build", 12),
    ("test_files", "Automated Test File", 48),
]
JOURNEY_COLUMNS = [
    ("key", "Journey Step", 11),
    ("journey", "Journey", 34),
    ("step", "Step", 70),
    ("result", "Result", 10),
    ("actual", "Actual Result", 40),
    ("defect", "Defect ID", 12),
    ("tester", "Tester", 12),
    ("test_date", "Test Date", 12),
    ("build", "Build", 12),
]
DEFECT_COLUMNS = [
    ("Defect ID", 11), ("Title", 44), ("Case or Journey Step", 18), ("Module", 14), ("Platform", 10), ("Role", 10),
    ("Severity", 10), ("Priority", 9), ("Steps to Reproduce", 50), ("Expected", 36), ("Actual", 36),
    ("Evidence", 24), ("Build", 12), ("Raised By", 12), ("Raised On", 12), ("Status", 12), ("Notes", 30),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RUN_FILLS = {"passed": "C6EFCE", "failed": "FFC7CE", "xfailed": "FFEB9C", "xpassed": "FFC7CE", "skipped": "D9D9D9"}
RESULT_FILLS = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFEB9C"}


def split_row(line: str) -> list:
    cells, current, in_code = [], [], False
    for char in line.strip().strip("|"):
        if char == "`":
            in_code = not in_code
        if char == "|" and not in_code:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    cells.append("".join(current).strip())
    return cells


def clean(text: str) -> str:
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.I)
    return text.replace("`", "").replace("\\|", "|").strip()


def read_specs() -> list:
    rows = []
    for path in sorted(glob.glob(os.path.join(SPECS, "*.md"))):
        name = os.path.splitext(os.path.basename(path))[0]
        if name.startswith("_") or name.upper() == "README":
            continue
        module = MODULE_TITLES.get(name, name)
        feature = ""
        with open(path, encoding="utf-8") as handle:
            for line in handle:
                feature_match = FEATURE_RE.match(line)
                if feature_match:
                    feature = f"{feature_match.group(1)} {feature_match.group(2).strip()}"
                    continue
                if not ROW_RE.match(line):
                    continue
                cells = [clean(cell) for cell in split_row(line)]
                case_id = cells[0]
                row = {
                    "id": case_id,
                    "module": module,
                    "feature": feature,
                    "phase": PHASES[case_id[-3]],
                    "priority": "",
                    "platform": "",
                    "role": "",
                    "preconditions": "",
                    "doc": os.path.basename(path),
                }
                if len(cells) >= 8:
                    row.update(
                        priority=cells[1], platform=cells[2], role=cells[3], preconditions=cells[4],
                        steps=cells[5], expected=cells[6], doc_status=cells[7],
                    )
                else:
                    row.update(
                        steps=cells[1] if len(cells) > 1 else "",
                        expected=cells[2] if len(cells) > 2 else "",
                        doc_status=cells[3] if len(cells) > 3 else "",
                    )
                    if row["phase"] == "UI":
                        lead = row["steps"].split(":", 1)[0].strip().lower()
                        row["platform"] = {"web": "Web", "mobile": "Mobile"}.get(lead, "")
                rows.append(row)
    return rows


def read_journeys() -> list:
    path = os.path.join(SPECS, "_journeys.md")
    if not os.path.isfile(path):
        return []
    rows, journey = [], None
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.rstrip()
            match = JOURNEY_RE.match(line)
            if match:
                journey = f"{match.group(1)} {match.group(2).strip()}"
                code = match.group(1)
                continue
            if not journey:
                continue
            if line.startswith("## "):
                journey = None
                continue
            step = STEP_RE.match(line)
            if step:
                rows.append({"key": f"{code}-S{int(step.group(1)):02d}", "journey": journey, "step": clean(step.group(2))})
            elif line.startswith("Exit checks:"):
                rows.append({"key": f"{code}-EXIT", "journey": journey, "step": clean(line)})
    return rows


def find_automation() -> dict:
    roots = [os.path.join(ROOT, "api_tests"), os.path.join(ROOT, "ui_tests", "tests")]
    app = os.environ.get("COS360_APP") or os.path.join(os.path.dirname(ROOT), "COS360_Full_App")
    for relative in ("backend/tests/unit", "web/src", "mobile/__tests__"):
        candidate = os.path.join(app, relative)
        if os.path.isdir(candidate):
            roots.append(candidate)
    found = defaultdict(set)
    for base in roots:
        for dirpath, dirnames, filenames in os.walk(base):
            dirnames[:] = [d for d in dirnames if d not in {"node_modules", "__pycache__"}]
            for filename in filenames:
                if not filename.endswith((".py", ".js", ".ts", ".tsx")):
                    continue
                if base.endswith("src") and ".test." not in filename:
                    continue
                path = os.path.join(dirpath, filename)
                with open(path, encoding="utf-8", errors="ignore") as handle:
                    text = handle.read()
                for case_id in set(ID_RE.findall(text)):
                    found[case_id].add(os.path.relpath(path, os.path.dirname(ROOT)).replace("\\", "/"))
    return found


def read_results() -> dict:
    merged = {}
    for path in glob.glob(os.path.join(REPORTS, "results-*.json")):
        with open(path, encoding="utf-8") as handle:
            for _, entry in json.load(handle).items():
                for case_id in entry["tc"]:
                    merged.setdefault(case_id, []).append(entry["outcome"])
    try:
        import ui_results
        for case_id, entries in ui_results.load().items():
            merged.setdefault(case_id, []).extend(outcome for outcome, _, _ in entries)
    except ImportError:
        pass
    summary = {}
    for case_id, outcomes in merged.items():
        for status in ("failed", "xpassed", "xfailed", "skipped", "passed"):
            if status in outcomes:
                summary[case_id] = status
                break
    return summary


def read_previous(sheet_name: str, key_title: str, fields: list) -> dict:
    """Carry tester-entered columns over from the existing workbook, keyed by ID."""
    if not os.path.isfile(OUT_XLSX):
        return {}
    try:
        wb = load_workbook(OUT_XLSX, read_only=True)
    except Exception:
        return {}
    if sheet_name not in wb.sheetnames:
        return {}
    rows = wb[sheet_name].iter_rows(values_only=True)
    header = list(next(rows, []) or [])
    if key_title not in header:
        return {}
    index = {title: header.index(title) for _, title, _ in fields if title in header}
    out = {}
    for row in rows:
        key = row[header.index(key_title)]
        if not key:
            continue
        values = {k: row[index[t]] for k, t, _ in fields if t in index and row[index[t]] not in (None, "")}
        if values:
            out[key] = values
    return out


def read_previous_defects() -> list:
    if not os.path.isfile(OUT_XLSX):
        return []
    try:
        wb = load_workbook(OUT_XLSX, read_only=True)
    except Exception:
        return []
    if "Defects" not in wb.sheetnames:
        return []
    rows = list(wb["Defects"].iter_rows(min_row=2, values_only=True))
    return [r for r in rows if any(v not in (None, "") for v in r)]


def write_header(sheet, columns) -> None:
    for index, column in enumerate(columns, start=1):
        title, width = (column[1], column[2]) if len(column) == 3 else column
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"


def write_cases(sheet, rows) -> None:
    write_header(sheet, CASE_COLUMNS)
    keys = [key for key, _, _ in CASE_COLUMNS]
    run_col = keys.index("last_run") + 1
    result_col = keys.index("result") + 1
    for r, row in enumerate(rows, start=2):
        for c, key in enumerate(keys, start=1):
            cell = sheet.cell(row=r, column=c, value=row.get(key, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        fill = RUN_FILLS.get(row["last_run"])
        if fill:
            sheet.cell(row=r, column=run_col).fill = PatternFill("solid", fgColor=fill)
        fill = RESULT_FILLS.get(row.get("result"))
        if fill:
            sheet.cell(row=r, column=result_col).fill = PatternFill("solid", fgColor=fill)
    last = max(len(rows) + 1, 2)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(keys))}{last}"
    validation = DataValidation(type="list", formula1=RESULT_VALUES, allow_blank=True)
    sheet.add_data_validation(validation)
    validation.add(f"{get_column_letter(result_col)}2:{get_column_letter(result_col)}{last}")


def ui_defects(rows: list) -> dict:
    """Known UI defects come from the test.fail reasons ("UI-XXX-NN: ...") in the latest UI results."""
    try:
        import ui_results
    except ImportError:
        return {}
    triage_path = os.path.join(ROOT, "tools", "defect_triage.json")
    triage = json.load(open(triage_path, encoding="utf-8")) if os.path.isfile(triage_path) else {}
    by_id = {row["id"]: row for row in rows}
    found = {}
    for case_id, entries in ui_results.load().items():
        for outcome, reason, project in entries:
            match = re.match(r"(UI-[A-Z]+-\d+)\s*:?\s*(.*)", reason or "")
            if outcome != "xfailed" or not match:
                continue
            item = found.setdefault(match.group(1), {"title": match.group(2), "cases": set(), "platforms": set()})
            item["cases"].add(case_id)
            item["platforms"].add(project.capitalize())
    out = {}
    for defect_id, item in found.items():
        cases = sorted(item["cases"])
        first = by_id.get(cases[0], {})
        info = triage.get(defect_id, {})
        out[defect_id] = [
            defect_id, item["title"][:200], ", ".join(cases), first.get("module", ""), "/".join(sorted(item["platforms"])),
            first.get("role", ""), info.get("severity", ""), info.get("priority", ""),
            "Steps are in the first listed test case.", "As the test case states.", info.get("note", ""),
            "Automated test (known defect)", "", "UI automation", "2026-10-07", "Open", "",
        ]
    return out


def main() -> None:
    rows = read_specs()
    if not rows:
        sys.exit(f"No test cases found under {SPECS}")
    journeys = read_journeys()
    automation = find_automation()
    results = read_results()
    previous = read_previous("Test Cases", "Test Case ID", [c for c in CASE_COLUMNS if c[0] in MANUAL_FIELDS])
    previous_journeys = read_previous("Journeys", "Journey Step", [c for c in JOURNEY_COLUMNS if c[0] in MANUAL_FIELDS])
    previous_defects = read_previous_defects()
    for row in rows:
        files = sorted(automation.get(row["id"], []))
        row["automated"] = "Yes" if files else "No"
        row["test_files"] = "\n".join(files)
        row["last_run"] = results.get(row["id"], "not run")
        row.update(previous.get(row["id"], {}))
    for row in journeys:
        row.update(previous_journeys.get(row["key"], {}))

    os.makedirs(os.path.dirname(OUT_XLSX), exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    write_cases(wb.create_sheet("Test Cases"), rows)
    smoke = [r for r in rows if r["phase"] == "UI" and r["priority"] == "P1"]
    write_cases(wb.create_sheet("Smoke"), smoke)

    jsheet = wb.create_sheet("Journeys")
    write_header(jsheet, JOURNEY_COLUMNS)
    jkeys = [key for key, _, _ in JOURNEY_COLUMNS]
    for r, row in enumerate(journeys, start=2):
        for c, key in enumerate(jkeys, start=1):
            cell = jsheet.cell(row=r, column=c, value=row.get(key, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if journeys:
        validation = DataValidation(type="list", formula1=RESULT_VALUES, allow_blank=True)
        jsheet.add_data_validation(validation)
        col = get_column_letter(jkeys.index("result") + 1)
        validation.add(f"{col}2:{col}{len(journeys) + 1}")
        jsheet.auto_filter.ref = f"A1:{get_column_letter(len(jkeys))}{len(journeys) + 1}"

    dsheet = wb.create_sheet("Defects")
    write_header(dsheet, DEFECT_COLUMNS)
    automatic = ui_defects(rows)
    for defect in previous_defects:
        if defect[0] not in automatic:
            dsheet.append(list(defect))
    for defect_id in sorted(automatic):
        dsheet.append(automatic[defect_id])

    summary.append(["COS360 test case catalog"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    header = ["Module", "Total", "Unit", "API", "UI", "UI P1", "UI P2", "UI P3", "Automated", "Passed", "Failed",
              "Known defect (xfail)", "Skipped", "Not run", "Manual Pass", "Manual Fail", "Manual Blocked"]
    summary.append(header)
    for cell in summary[3]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    by_module = defaultdict(list)
    for row in rows:
        by_module[row["module"]].append(row)
    totals = Counter()
    for module in MODULE_TITLES.values():
        items = by_module.get(module, [])
        if not items:
            continue
        phase = Counter(item["phase"] for item in items)
        prio = Counter(item["priority"] for item in items if item["phase"] == "UI")
        run = Counter(item["last_run"] for item in items)
        manual = Counter(item.get("result") for item in items)
        line = [
            module, len(items), phase["Unit"], phase["API"], phase["UI"], prio["P1"], prio["P2"], prio["P3"],
            sum(1 for item in items if item["automated"] == "Yes"),
            run["passed"], run["failed"] + run["xpassed"], run["xfailed"], run["skipped"], run["not run"],
            manual["Pass"], manual["Fail"], manual["Blocked"],
        ]
        summary.append(line)
        for index, value in enumerate(line[1:], start=1):
            totals[index] += value
    summary.append(["TOTAL"] + [totals[i] for i in range(1, len(header))])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 26
    for index in range(2, len(header) + 1):
        summary.column_dimensions[get_column_letter(index)].width = 12
    summary.append([])
    summary.append([f"Journeys: {len({r['journey'] for r in journeys})} journeys, {len(journeys)} steps (sheet Journeys)."])
    summary.append(["Types: Unit (backend, web, mobile unit tests), API (pytest against the QA tenant), UI (manual cases and Playwright)."])
    summary.append(["Last Run comes from reports/results-*.json of the latest automated run. Result, Actual Result, Defect ID, Tester, Test Date and Build are filled by manual testers and survive a rebuild."])
    summary.append(["How to test by hand: test_specs/_manual-testing-guide.md. Smoke sheet = every P1 UI case."])

    wb.save(OUT_XLSX)
    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow([title for _, title, _ in CASE_COLUMNS])
        for row in rows:
            writer.writerow([str(row.get(key, "") or "").replace("\n", "; ") for key, _, _ in CASE_COLUMNS])

    print(f"{len(rows)} test cases ({len(smoke)} smoke), {len(journeys)} journey steps -> {os.path.relpath(OUT_XLSX, ROOT)}")
    print(f"automated: {sum(1 for r in rows if r['automated'] == 'Yes')}, with a last run result: {sum(1 for r in rows if r['last_run'] != 'not run')}")


if __name__ == "__main__":
    main()
