import io
import uuid

import pytest
from openpyxl import load_workbook

from api_tests.reports.helpers import ROLES, make_class, make_student
from api_tests.support import Cleanup

MODULE_NODES = {"Masters", "Students", "Staff", "Fee", "Transport", "Exam", "Expense", "Reports", "Administration"}


def _names(menu):
    return {node["name"] for node in menu}


def _flat(menu):
    result = []
    for node in menu:
        result.append(node)
        result.extend(_flat(node.get("children") or []))
    return result


@pytest.mark.api
@pytest.mark.tc("TC-RPT-01-A01")
def test_admin_login_menu_and_permissions(logins):
    payload = logins["admin"]
    assert MODULE_NODES <= _names(payload["menu"])
    assert isinstance(payload["permissions"], dict) and payload["permissions"]


@pytest.mark.api
@pytest.mark.tc("TC-RPT-01-A02")
def test_teacher_has_no_fee_grants(logins):
    permissions = logins["teacher"]["permissions"]
    assert not [key for key in permissions if key.startswith("fee_")]


@pytest.mark.api
@pytest.mark.tc("TC-RPT-01-A03")
@pytest.mark.parametrize("role", ["student", "parent"])
def test_student_and_parent_menu_allowlist(logins, role):
    names = _names(logins[role]["menu"])
    assert names <= {"Dashboard", "Students", "Fee", "Exam"}
    assert "Dashboard" in names


@pytest.mark.api
@pytest.mark.tc("TC-RPT-02-A01")
def test_admin_menu_nodes_have_children_and_paths(logins):
    nodes = {node["name"]: node for node in logins["admin"]["menu"]}
    for name in ("Masters", "Students", "Transport", "Expense", "Fee"):
        assert nodes[name]["children"], name
        assert all(child.get("path") for child in nodes[name]["children"])
    assert nodes["Transport"]["path"] == "/transport"


@pytest.mark.api
@pytest.mark.tc("TC-RPT-02-A02")
def test_student_menu_is_subset_of_admin_menu(logins):
    admin_paths = {node.get("path") for node in _flat(logins["admin"]["menu"])}
    student_paths = {node.get("path") for node in _flat(logins["student"]["menu"])}
    assert student_paths <= admin_paths
    assert len(student_paths) < len(admin_paths)


@pytest.mark.api
@pytest.mark.tc("TC-RPT-03-A01")
def test_admin_report_permissions(logins):
    permissions = logins["admin"]["permissions"]
    assert "read" in permissions["fee_reports"]
    assert "read" in permissions["expense_reports"]


@pytest.mark.api
@pytest.mark.tc("TC-RPT-03-A02")
def test_teacher_report_permissions(logins):
    permissions = logins["teacher"]["permissions"]
    assert "student_reports" in permissions and "attendance_reports" in permissions
    assert "fee_reports" not in permissions


@pytest.mark.api
@pytest.mark.tc("TC-RPT-10-A01")
def test_no_academic_report_endpoint(admin):
    assert admin.get("/reports/academic").status_code == 404


@pytest.mark.api
@pytest.mark.tc("TC-RPT-10-A02")
@pytest.mark.parametrize("role", ["admin", "teacher"])
def test_exam_list_for_academic_screen(role_clients, role):
    assert role_clients[role].get("/exams").status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-RPT-11-A01")
def test_no_transport_report_endpoint(admin):
    assert admin.get("/reports/transport").status_code == 404


@pytest.mark.api
@pytest.mark.tc("TC-RPT-11-A02")
def test_teacher_reads_transport_lists(teacher):
    for path in ("/masters/routes/all_routes", "/masters/vehicles/", "/masters/trips/"):
        assert teacher.get(path).status_code == 200, path


@pytest.fixture(scope="module")
def sample(admin, academic_year_id):
    stack = Cleanup()
    try:
        klass = make_class(admin, stack, academic_year_id, ("A",))
        make_student(admin, stack, academic_year_id, klass, 0, "Pune")
        yield {"class_id": klass["id"]}
    finally:
        stack.run()


def _exports(sample):
    return [
        ("/reports/students/export", {"report_type": "student_summary", "filters": {"class_id": sample["class_id"]}}),
        ("/reports/staff/export", {"report_type": "staff_summary", "filters": {"gender": "Male"}}),
        ("/reports/fees/export", {"report_type": "fee_structure", "filters": {}}),
        ("/reports/attendance/export", {"report_type": "student_attendance", "filters": {}}),
        ("/reports/financial/export", {"report_type": "ledger", "filters": {}}),
    ]


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A01")
def test_list_envelope_shape(admin, sample):
    for path in (
        "/reports/students/summary",
        "/reports/staff/summary",
        "/reports/fees/collection-summary",
        "/reports/fees/pending-fees",
        "/reports/fees/fee-structure",
        "/reports/attendance/students",
        "/reports/attendance/staff",
        "/reports/financial/expenditure",
        "/reports/financial/ledger",
    ):
        response = admin.get(path, params={"page_size": 1})
        assert response.status_code == 200, path
        body = response.json()
        assert {"data", "total_count", "page", "page_size", "total_pages"} <= set(body), path
        for key in ("total_count", "page", "page_size", "total_pages"):
            assert isinstance(body[key], int), (path, key)


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A02")
def test_csv_exports_for_every_group(admin, sample):
    for path, base in _exports(sample):
        response = admin.post(path, json={**base, "format": "csv"})
        assert response.status_code == 200, path
        assert response.headers["content-type"].startswith("text/csv"), path
        assert "attachment" in response.headers["content-disposition"], path
        assert int(response.headers["content-length"]) == len(response.content), path


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A03")
def test_xlsx_export_is_a_workbook(admin, sample):
    response = admin.post(
        "/reports/students/export",
        json={"report_type": "student_summary", "filters": {"class_id": sample["class_id"]}, "format": "xlsx"},
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert header == ["sl_no", "admission_no", "student_id", "class_section", "address", "city", "academic_year"]
    assert sheet.max_row == 2


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A04")
def test_pdf_export(admin, sample):
    response = admin.post(
        "/reports/students/export",
        json={"report_type": "student_summary", "filters": {"class_id": sample["class_id"]}, "format": "pdf"},
    )
    assert response.status_code == 200 and response.content.startswith(b"%PDF")


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A05")
def test_export_unsupported_format(admin):
    response = admin.post("/reports/students/export", json={"report_type": "student_summary", "filters": {}, "format": "xls"})
    assert response.status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A06")
def test_export_filename_is_not_sanitised(admin, sample):
    response = admin.post(
        "/reports/students/export",
        json={
            "report_type": "student_summary",
            "filters": {"class_id": sample["class_id"]},
            "format": "csv",
            "filename": "qa report.v1",
        },
    )
    assert response.headers["content-disposition"] == "attachment; filename=qa report.v1.csv"


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A07")
@pytest.mark.skip(reason="creating 1500 rows is impractical and they could not be cleaned up afterwards")
def test_export_of_1500_rows():
    pass


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A08")
def test_export_body_validation(admin):
    assert admin.post("/reports/students/export", json={"report_type": "student_summary", "format": "csv"}).status_code == 422
    assert admin.post("/reports/students/export", json={"filters": {}, "format": "csv"}).status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-RPT-12-A09")
def test_unauthenticated(anon):
    assert anon.get("/reports/fees/pending-fees").status_code == 401
    assert anon.post("/reports/fees/export", json={"report_type": "fee_structure", "filters": {}, "format": "csv"}).status_code == 401


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A01")
def test_exports_are_not_audited(admin, sample):
    export = admin.post(
        "/reports/students/export",
        json={"report_type": "student_summary", "filters": {"class_id": sample["class_id"]}, "format": "csv"},
    )
    assert export.status_code == 200
    response = admin.get("/reports/audit")
    assert response.status_code == 200
    assert response.json() == []


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A02")
@pytest.mark.parametrize("params", [{"page_size": 101}, {"page": 0}, {"page_size": 0}])
def test_audit_paging_bounds(admin, params):
    assert admin.get("/reports/audit", params=params).status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A03")
def test_audit_record_lookup(admin):
    missing = admin.get(f"/reports/audit/{uuid.uuid4()}")
    assert missing.status_code == 404 and missing.json()["detail"] == "Export record not found"
    assert admin.get("/reports/audit/not-a-uuid").status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A04")
def test_download_unknown_record(admin):
    response = admin.get(f"/reports/download/{uuid.uuid4()}")
    assert response.status_code == 404 and response.json()["detail"] == "Export record not found"


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A05")
@pytest.mark.skip(reason="needs a completed report_audit row with a file on disk; no endpoint writes the table and tests may not touch the database")
def test_download_completed_export():
    pass


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A06")
@pytest.mark.skip(reason="needs a pending report_audit row; no endpoint writes the table and tests may not touch the database")
def test_download_pending_export():
    pass


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A07")
def test_foreign_record_ids_are_unknown(admin, tenant_b):
    foreign = uuid.uuid4()
    assert admin.get(f"/reports/audit/{foreign}").status_code == 404
    assert tenant_b.get(f"/reports/audit/{foreign}").status_code == 404
    assert tenant_b.get(f"/reports/download/{foreign}").status_code == 404


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A08")
@pytest.mark.parametrize("role", ROLES)
def test_audit_role_matrix(role_clients, role):
    client = role_clients[role]
    allowed = role in ("admin", "staff", "teacher")
    assert client.get(f"/reports/audit/{uuid.uuid4()}").status_code == (404 if allowed else 403)
    assert client.get(f"/reports/download/{uuid.uuid4()}").status_code == (404 if allowed else 403)
    listing = client.get("/reports/audit")
    assert listing.status_code == 200 if allowed else listing.status_code >= 400


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A08")
@pytest.mark.parametrize("role", ["student", "parent"])
def test_audit_list_denied_is_403(role_clients, role):
    assert role_clients[role].get("/reports/audit").status_code == 403


@pytest.mark.api
@pytest.mark.tc("TC-RPT-13-A09")
def test_audit_unauthenticated(anon):
    assert anon.get("/reports/audit").status_code == 401
    assert anon.get(f"/reports/audit/{uuid.uuid4()}").status_code == 401
    assert anon.get(f"/reports/download/{uuid.uuid4()}").status_code == 401
