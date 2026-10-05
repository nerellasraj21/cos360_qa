import io

import openpyxl
import pytest

from api_tests.students import helpers as h
from api_tests.support import unique

UPLOAD = "/students/admission/bulk-upload"
TEMPLATE = "/students/admission/bulk-upload/template"
EMPTY_SHEET = None


def send(client, content, name="rows.xlsx", ctype=h.XLSX_TYPE, **kwargs):
    return client.post(UPLOAD, files={"file": (name, io.BytesIO(content), ctype)}, **kwargs)


def register(admin, klass, cleanup):
    listing = admin.get("/students/admission/", params={"class_id": klass["id"], "limit": 100})
    for item in listing.json()["items"]:
        cleanup.delete_later(admin, f"/students/admission/{item['id']}")


def upload_rows(admin, klass, rows, cleanup, **kwargs):
    response = send(admin, h.build_workbook(rows), **kwargs)
    register(admin, klass, cleanup)
    return response


def numbered(klass, count, **overrides):
    return [h.bulk_row(klass, **{"Admission no": unique("stu_bn"), **overrides}) for _ in range(count)]


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A01")
def test_three_valid_rows_are_created(admin, klass, cleanup):
    rows = numbered(klass, 3)
    response = upload_rows(admin, klass, rows, cleanup)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["errors"] == []
    assert body["total_rows"] == 3
    assert [c["row"] for c in body["created"]] == [2, 3, 4]
    assert [c["admission_number"] for c in body["created"]] == [r["Admission no"] for r in rows]
    assert set(body["created"][0]) == {"row", "student_id", "admission_number", "name"}
    assert admin.get(f"/students/admission/id/{body['created'][0]['student_id']}").status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A02")
def test_bad_class_row_is_reported_and_others_created(admin, klass, cleanup):
    rows = numbered(klass, 3)
    rows[1]["joining class"] = klass["name"][:-1] + "Z"
    response = upload_rows(admin, klass, rows, cleanup)
    assert response.status_code == 200
    body = response.json()
    assert len(body["created"]) == 2
    assert body["total_rows"] == 3
    assert len(body["errors"]) == 1
    assert body["errors"][0].startswith(f"Row 3: Class '{rows[1]['joining class']}' not found")
    assert f"did you mean '{klass['name']}'" in body["errors"][0]


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A03")
def test_missing_required_header(admin):
    headers = [name for name in h.BULK_HEADERS if name != "Father phone"]
    response = send(admin, h.build_workbook([], headers))
    assert response.status_code == 400
    assert response.json()["detail"] == "Missing required column(s): Father phone"


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A04")
def test_non_excel_and_corrupt_files(admin):
    csv = send(admin, b"a,b\n1,2\n", name="rows.csv", ctype="text/csv")
    assert csv.status_code == 400
    assert csv.json()["detail"] == "File must be an Excel (.xlsx/.xls) file"
    broken = send(admin, b"this is not a workbook")
    assert broken.status_code == 400
    assert broken.json()["detail"].startswith("Invalid Excel file")


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A05")
def test_row_level_errors(admin, klass, cleanup):
    dup_number = unique("stu_bd")
    shared_email = f"{unique('stu_be')}@example.com"
    rows = [
        h.bulk_row(klass, **{"First name": None, "Admission no": unique("stu_bn")}),
        h.bulk_row(klass, **{"Father phone": None, "Admission no": unique("stu_bn")}),
        h.bulk_row(klass, **{"Address Line 1": None, "Admission no": unique("stu_bn")}),
        h.bulk_row(klass, **{"Date of birth (DD-MM-YYYY)": "not-a-date", "Admission no": unique("stu_bn")}),
        h.bulk_row(klass, **{"Aadhar No": "12345678901", "Admission no": unique("stu_bn")}),
        h.bulk_row(klass, **{"Admission no": dup_number}),
        h.bulk_row(klass, **{"Admission no": dup_number}),
        h.bulk_row(
            klass,
            **{"Father Email": shared_email, "Mother Email": shared_email, "Admission no": unique("stu_bn")},
        ),
    ]
    response = upload_rows(admin, klass, rows, cleanup)
    assert response.status_code == 200
    body = response.json()
    errors = {int(e.split(":")[0].split()[1]): e for e in body["errors"]}
    assert "First name is required" in errors[2]
    assert "Father name and Father phone are required" in errors[3]
    assert "Address Line 1 is required" in errors[4]
    assert "Date of birth 'not-a-date' is not a valid date" in errors[5]
    assert "12-digit" in errors[6]
    assert f"Admission number '{dup_number}' is already in use" in errors[8]
    assert "cannot have the same email" in errors[9]
    assert [c["row"] for c in body["created"]] == [7]
    assert body["total_rows"] == 8


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A06")
@pytest.mark.skip(reason="blocked: needs a tenant without an active academic year, and the active year is a tenant-wide singleton")
def test_no_active_academic_year():
    raise AssertionError


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A07")
def test_pre_primary_student_type_gets_year_number(admin, klass, cleanup):
    body = None
    for _ in range(4):
        row = h.bulk_row(klass, **{"Student type": "Pre Primary", "Admission no": None})
        response = upload_rows(admin, klass, [row], cleanup)
        assert response.status_code == 200
        body = response.json()
        if body["created"]:
            break
    assert body["created"], body
    assert body["created"][0]["admission_number"].startswith(str(h.today().year))


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A08")
def test_empty_rows_are_skipped(admin, klass, cleanup):
    rows = numbered(klass, 1) + [{}] + numbered(klass, 1)
    response = upload_rows(admin, klass, rows, cleanup)
    body = response.json()
    assert response.status_code == 200
    assert body["total_rows"] == 2
    assert len(body["created"]) == 2 and body["errors"] == []


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A09")
def test_blank_template(admin, klass):
    response = admin.get(TEMPLATE, params={"include_data": "false"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(h.XLSX_TYPE)
    assert "student_admission_bulk_upload_template.xlsx" in response.headers["content-disposition"]
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[:2] == ["Student Admission", "Lists"]
    assert book["Lists"].sheet_state == "hidden"
    listed = {cell.value for row in book["Lists"].iter_rows() for cell in row if cell.value}
    assert klass["name"] in listed


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A10")
def test_prefilled_template_contains_admissions(admin, family):
    response = admin.get(TEMPLATE, params={"include_data": "true"})
    assert response.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(response.content))["Student Admission"]
    header = [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]
    number_col = header.index("Admission no")
    dob_col = next(i for i, name in enumerate(header) if name.lower().startswith("date of birth"))
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
    mine = [r for r in rows if r[number_col] == family.other.number]
    assert mine, "prefilled template has no row for the new admission"
    assert mine[0][dob_col] == "10-04-2015"


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A11")
@pytest.mark.parametrize(
    "role,status", [("admin", 200), ("staff", 200), ("teacher", 403), ("student", 403), ("parent", 403)]
)
def test_bulk_permission_matrix(role_clients, role, status):
    empty = h.build_workbook([])
    assert send(role_clients[role], empty).status_code == status
    assert role_clients[role].get(TEMPLATE).status_code == status
    assert role_clients[role].get(TEMPLATE, params={"include_data": "true"}).status_code == status


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A12")
def test_bulk_requires_token(anon):
    assert send(anon, h.build_workbook([])).status_code == 401
    assert anon.get(TEMPLATE).status_code == 401


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A13")
def test_class_from_other_tenant_is_not_found(admin, tenant_b, tenant_b_name, klass):
    rows = [h.bulk_row(klass, **{"Admission no": unique("stu_bn")})]
    response = send(tenant_b, h.build_workbook(rows))
    assert response.status_code == 200
    body = response.json()
    assert body["created"] == []
    assert f"Class '{klass['name']}' not found" in body["errors"][0]
    mismatch = send(admin, h.build_workbook(rows), headers={"cschema": tenant_b_name})
    assert mismatch.status_code == 403


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A14")
def test_reupload_fails_on_duplicate_numbers(admin, klass, cleanup):
    rows = numbered(klass, 2)
    first = upload_rows(admin, klass, rows, cleanup)
    assert len(first.json()["created"]) == 2
    second = upload_rows(admin, klass, rows, cleanup)
    body = second.json()
    assert body["created"] == []
    assert len(body["errors"]) == 2
    assert all("is already in use" in e for e in body["errors"])


@pytest.mark.api
@pytest.mark.tc("TC-STU-10-A02")
@pytest.mark.xfail(
    strict=True,
    reason="STU-BUG-3: bulk upload resolves 'joining section' by name across the whole tenant, so a section name used by more than one class fails the row",
)
def test_section_is_resolved_within_the_class(admin, klass, cleanup):
    rows = numbered(klass, 1, **{"joining section": "A"})
    response = upload_rows(admin, klass, rows, cleanup)
    body = response.json()
    assert body["errors"] == [], body
    assert len(body["created"]) == 1
    detail = admin.get(f"/students/admission/id/{body['created'][0]['student_id']}").json()
    assert detail["admitted_section_id"] == klass["sections"]["A"]
