import contextlib
import io
import os
import tempfile
import time
import uuid
from datetime import date, timedelta

import pytest

from api_tests.support import Api, QA_TENANT, items_of, unique

NIL_UUID = "00000000-0000-0000-0000-000000000000"
OTHER_ROLES = ["teacher", "staff", "student", "parent"]
ALL_ROLES = ["admin", "teacher", "staff", "student", "parent"]
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

GS1_BANDS = [
    (90, 100, "A+", 4.0, True),
    (80, 89.99, "A", 3.5, True),
    (70, 79.99, "B", 3.0, True),
    (60, 69.99, "C", 2.0, True),
    (35, 59.99, "D", 1.0, True),
    (0, 34.99, "F", 0.0, False),
]

DEFAULT_COMPONENTS = {
    "math": [("Written", 80), ("Oral", 20)],
    "sci": [("Written", 100)],
    "eng": [("Written", 100)],
}

S1_MARKS = {"math": (62, 18), "sci": (90,), "eng": (70,)}
S2_MARKS = {"math": (68, 17), "sci": (95,), "eng": (90,)}
S3_MARKS = {"math": (45, 10), "sci": (45,), "eng": (30,)}


def bands(spec=GS1_BANDS):
    return [
        {
            "from_percent": a,
            "to_percent": b,
            "grade_label": g,
            "gpa": gpa,
            "remarks": g,
            "is_pass": p,
            "sort_order": i,
        }
        for i, (a, b, g, gpa, p) in enumerate(spec)
    ]


def create_exam_scheme(admin, cleanup, name=None, spec=GS1_BANDS, default=False):
    name = name or unique("exm_gs_")
    response = admin.post("/grade-schemes/exam", json={"name": name, "description": "d", "is_default": default, "bands": bands(spec)})
    assert response.status_code == 201, response.text
    data = response.json()
    cleanup.delete_later(admin, f"/grade-schemes/exam/{data['id']}")
    return data


def create_subject_scheme(admin, cleanup, name=None, spec=GS1_BANDS):
    name = name or unique("exm_ss_")
    response = admin.post("/grade-schemes/subject", json={"name": name, "description": "d", "is_default": False, "bands": bands(spec)})
    assert response.status_code == 201, response.text
    data = response.json()
    cleanup.delete_later(admin, f"/grade-schemes/subject/{data['id']}")
    return data


def create_remark_set(admin, cleanup, name=None):
    name = name or unique("exm_rs_")
    response = admin.post(
        "/remark-grades",
        json={"name": name, "options": [{"grade_letter": "A", "label": "Excellent", "sort_order": 0}, {"grade_letter": "B", "label": "Good", "sort_order": 1}]},
    )
    assert response.status_code == 201, response.text
    data = response.json()
    cleanup.delete_later(admin, f"/remark-grades/{data['id']}")
    return data


def remove_exam(admin, exam_id):
    response = admin.delete(f"/exams/{exam_id}")
    if response.status_code == 409:
        admin.post(f"/exams/{exam_id}/unlock", json={"reason": "cleanup"})
        admin.delete(f"/exams/{exam_id}")


def exam_payload(w, name=None, sections=("a",), subjects=("math", "sci", "eng"), comps=None, dates=False, scheme_id=None, **exam_extra):
    comps = comps or DEFAULT_COMPONENTS
    class_sections = []
    configs = []
    date_rows = []
    for sec in sections:
        sid = w.sections[sec]
        class_sections.append({"class_id": w.class_id, "section_id": sid})
        for idx, sub in enumerate(subjects):
            configs.append(
                {
                    "class_id": w.class_id,
                    "section_id": sid,
                    "subject_id": w.subjects[sub],
                    "sort_order": idx + 1,
                    "components": [
                        {
                            "component_name": cn,
                            "entry_type": "marks",
                            "max_marks": mx,
                            "include_in_total": True,
                            "sort_order": ci + 1,
                        }
                        for ci, (cn, mx) in enumerate(comps[sub])
                    ],
                }
            )
            if dates:
                date_rows.append(
                    {
                        "class_id": w.class_id,
                        "section_id": sid,
                        "subject_id": w.subjects[sub],
                        "exam_date": (date(2027, 2, 1) + timedelta(days=idx)).isoformat(),
                        "start_time": "09:30",
                        "end_time": "11:30",
                        "venue": "Hall A",
                    }
                )
    exam = {
        "exam_name": name or unique("exm_ex_"),
        "board": "State",
        "level": "primary",
        "exam_type": "Unit Test",
        "nature": "formative",
        "academic_year_id": w.year_id,
        "exam_grade_scheme_id": scheme_id or w.scheme_id,
    }
    exam.update(exam_extra)
    return {"exam": exam, "class_sections": class_sections, "subject_configs": configs, "exam_dates": date_rows}


def create_exam(admin, cleanup, w, **kwargs):
    payload = exam_payload(w, **kwargs)
    response = admin.post("/exams", json=payload)
    assert response.status_code == 201, response.text
    exam_id = response.json()["exam_id"]
    cleanup.add(remove_exam, admin, exam_id)
    return exam_id


def get_configs(admin, exam_id):
    response = admin.get(f"/exams/{exam_id}/subject-configs")
    assert response.status_code == 200, response.text
    return response.json()


def config_map(w, configs, section="a"):
    sid = w.sections[section]
    by_subject = {}
    inverse = {v: k for k, v in w.subjects.items()}
    for c in configs:
        if c["section_id"] == sid:
            by_subject[inverse[c["subject_id"]]] = c
    return by_subject


def comp_by_name(config, name):
    return next(c for c in config["components"] if c["component_name"] == name)


def save_marks(client, exam_id, config, rows, attempt=1):
    body = {
        "exam_id": exam_id,
        "subject_config_id": config["id"],
        "marks": rows,
        "attempt_number": attempt,
    }
    return client.post(f"/exams/{exam_id}/marks", json=body)


def mark_row(student_id, component, value=None, absent=False, remark=None):
    return {
        "student_id": student_id,
        "component_id": component["id"],
        "marks_obtained": value,
        "remark_grade": remark,
        "is_absent": absent,
    }


def enter_student_marks(admin, exam_id, w, configs, student, marks, section="a"):
    cmap = config_map(w, configs, section)
    for sub, values in marks.items():
        cfg = cmap[sub]
        comps = sorted(cfg["components"], key=lambda c: c["sort_order"])
        rows = []
        for comp, value in zip(comps, values):
            if value is None:
                continue
            rows.append(mark_row(student["id"], comp, value))
        if rows:
            response = save_marks(admin, exam_id, cfg, rows)
            assert response.status_code == 200, response.text


def standard_marks(admin, exam_id, w, configs, count=3):
    sets = [S1_MARKS, S2_MARKS, S3_MARKS]
    for student, marks in zip(w.students[:count], sets[:count]):
        enter_student_marks(admin, exam_id, w, configs, student, marks)


def grid(client, exam_id, w, config, section="a", **extra):
    params = {"class_id": w.class_id, "section_id": w.sections[section], "subject_config_id": config["id"]}
    params.update(extra)
    return client.get(f"/exams/{exam_id}/marks", params=params)


def audit_actions(admin, exam_id):
    response = admin.get(f"/exams/{exam_id}/audit", params={"page_size": 100})
    assert response.status_code == 200, response.text
    return response.json()


def build_xlsx(headers, rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_xlsx(content):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def other_tenant_header(api_token):
    client = Api(token=api_token, tenant_header="qa_school_b")
    return client


def new_uuid():
    return str(uuid.uuid4())


def denied(role_clients, roles, method, path, **kwargs):
    out = {}
    for role in roles:
        out[role] = getattr(role_clients[role], method)(path, **kwargs).status_code
    return out


@contextlib.contextmanager
def settings_lock(timeout=600):
    path = os.path.join(tempfile.gettempdir(), "cos360_exm_settings.lock")
    deadline = time.time() + timeout
    fd = None
    while fd is None:
        try:
            fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            try:
                if time.time() - os.path.getmtime(path) > 120:
                    os.remove(path)
                    continue
            except OSError:
                pass
            if time.time() > deadline:
                raise
            time.sleep(0.2)
    try:
        yield
    finally:
        os.close(fd)
        try:
            os.remove(path)
        except OSError:
            pass
