import io

import pytest
from openpyxl import Workbook, load_workbook

from api_tests.exam.helpers import (
    NIL_UUID,
    OTHER_ROLES,
    XLSX_TYPE,
    audit_actions,
    build_xlsx,
    comp_by_name,
    config_map,
    create_exam,
    create_remark_set,
    exam_payload,
    get_configs,
    grid,
    mark_row,
    new_uuid,
    read_xlsx,
    remove_exam,
    save_marks,
)
from api_tests.exam.conftest import World


def math_cfg(world, ex1):
    return config_map(world, ex1["configs"])["math"]


def written(cfg):
    return comp_by_name(cfg, "Written")


def oral(cfg):
    return comp_by_name(cfg, "Oral")


def student_row(rows, student):
    return next(r for r in rows if r["student_id"] == student["id"])


def tmpl_params(world, cfg, section="a"):
    return {"class_id": world.class_id, "section_id": world.sections[section], "subject_config_id": cfg["id"]}


def get_template(client, exam_id, world, cfg):
    return client.get(f"/exams/{exam_id}/marks/template", params=tmpl_params(world, cfg))


def upload(client, exam_id, world, cfg, content, name="m.xlsx", params=None, ctype=XLSX_TYPE):
    return client.post(
        f"/exams/{exam_id}/marks/upload",
        params=params if params is not None else tmpl_params(world, cfg),
        files={"file": (name, content, ctype)},
    )


def fill_template(content, values):
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        sid = row[0].value
        if sid in values:
            for offset, v in enumerate(values[sid]):
                row[3 + offset].value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A01")
def test_grid_lists_enrolled_students(teacher, world, ex1):
    cfg = math_cfg(world, ex1)
    response = grid(teacher, ex1["id"], world, cfg)
    assert response.status_code == 200, response.text
    rows = response.json()
    assert [r["student_id"] for r in rows] == [s["id"] for s in world.students]
    assert [r["admission_number"] for r in rows] == [s["admission_number"] for s in world.students]
    for r in rows:
        assert r["student_name"]
        for c in cfg["components"]:
            entry = r["marks"].get(c["id"])
            if entry is not None:
                assert entry["mark_id"] is None and entry["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A02")
def test_grid_with_nil_section(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    response = admin.get(
        f"/exams/{ex1['id']}/marks", params={"class_id": world.class_id, "section_id": NIL_UUID, "subject_config_id": cfg["id"]}
    )
    assert response.status_code == 200
    assert {r["student_id"] for r in response.json()} >= {s["id"] for s in world.students}


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A03")
@pytest.mark.parametrize("missing", ["class_id", "section_id", "subject_config_id"])
def test_grid_missing_params(admin, world, ex1, missing):
    cfg = math_cfg(world, ex1)
    params = tmpl_params(world, cfg)
    del params[missing]
    assert admin.get(f"/exams/{ex1['id']}/marks", params=params).status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A04")
def test_grid_pagination(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    response = grid(admin, ex1["id"], world, cfg, page_size=1, page=2)
    assert response.status_code == 200
    assert [r["student_id"] for r in response.json()] == [world.students[1]["id"]]


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A05")
def test_grid_unknown_config(admin, world, ex1):
    response = admin.get(
        f"/exams/{ex1['id']}/marks",
        params={"class_id": world.class_id, "section_id": world.sections["a"], "subject_config_id": new_uuid()},
    )
    assert response.status_code == 200
    rows = response.json()
    assert len(rows) == 3 and all(r["marks"] == {} for r in rows)


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A06")
@pytest.mark.parametrize("role", ["admin", "teacher", "staff"])
def test_grid_readers(role_clients, world, ex1, role):
    assert grid(role_clients[role], ex1["id"], world, math_cfg(world, ex1)).status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A07")
@pytest.mark.parametrize("role", ["student", "parent"])
def test_grid_denied(role_clients, world, ex1, role):
    assert grid(role_clients[role], ex1["id"], world, math_cfg(world, ex1)).status_code == 403


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A08")
def test_save_one_mark_as_teacher(teacher, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    response = save_marks(teacher, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 62)])
    assert response.status_code == 200, response.text
    rows = response.json()
    assert len(rows) == 1
    assert rows[0]["entry_source"] == "manual"
    assert rows[0]["marks_obtained"] == "62.00"
    assert rows[0]["student_id"] == s["id"] and rows[0]["component_id"] == written(cfg)["id"]
    cell = student_row(grid(teacher, ex1["id"], world, cfg).json(), s)["marks"][written(cfg)["id"]]
    assert cell["marks_obtained"] == 62.0


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A09")
def test_save_bulk(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    rows = []
    for i, s in enumerate(world.students):
        rows.append(mark_row(s["id"], written(cfg), 50 + i))
        rows.append(mark_row(s["id"], oral(cfg), 10 + i))
    response = save_marks(admin, ex1["id"], cfg, rows)
    assert response.status_code == 200
    assert len(response.json()) == 6


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A10")
def test_resave_updates_in_place(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    first = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 40)]).json()[0]
    second = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 45)]).json()
    row = next(r for r in second if r["student_id"] == s["id"])
    assert row["id"] == first["id"]
    assert row["marks_obtained"] == "45.00"
    assert row["updated_at"] is not None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A11")
def test_marks_equal_to_max(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    assert save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 80)]).status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A12")
@pytest.mark.parametrize("value", [80.01, 81])
def test_marks_above_max_rejected(admin, world, ex1, value):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    response = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), value)])
    assert response.status_code == 422
    assert "exceeds the component maximum" in response.json()["detail"]
    cell = student_row(grid(admin, ex1["id"], world, cfg).json(), s)["marks"].get(written(cfg)["id"])
    assert cell is None or cell["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A14")
def test_marks_zero(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    response = save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 0)])
    assert response.status_code == 200
    assert response.json()[0]["marks_obtained"] == "0.00"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A15")
def test_negative_marks_rejected(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    assert save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), -1)]).status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A16")
def test_null_clears_mark(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 33)])
    response = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), None)])
    assert response.status_code == 200
    assert next(r for r in response.json() if r["student_id"] == s["id"])["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A17")
def test_absent_flag(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    response = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), None, absent=True)])
    assert response.status_code == 200
    assert response.json()[0]["is_absent"] is True
    cell = student_row(grid(admin, ex1["id"], world, cfg).json(), s)["marks"][written(cfg)["id"]]
    assert cell["is_absent"] is True


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A18")
def test_absent_with_marks_stores_both(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    response = save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 40, absent=True)])
    assert response.status_code == 200
    assert response.json()[0]["is_absent"] is True
    assert response.json()[0]["marks_obtained"] == "40.00"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A19")
def test_marks_rounded_to_two_decimals(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    response = save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 45.678)])
    assert response.status_code == 200
    assert response.json()[0]["marks_obtained"] == "45.68"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A20")
def test_batch_with_one_invalid_saves_nothing(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    rows = [
        mark_row(world.students[0]["id"], written(cfg), 10),
        mark_row(world.students[1]["id"], written(cfg), 20),
        mark_row(world.students[2]["id"], written(cfg), 99),
    ]
    assert save_marks(admin, ex1["id"], cfg, rows).status_code == 422
    for r in grid(admin, ex1["id"], world, cfg).json():
        cell = r["marks"].get(written(cfg)["id"])
        assert cell is None or cell["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A21")
def test_save_payload_validation(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    assert save_marks(admin, ex1["id"], cfg, []).status_code == 422
    assert save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 5)], attempt=0).status_code == 422
    assert save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 5, remark="ABCDEF")]).status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A22")
def test_unknown_component_or_student_conflicts(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    bad_comp = {**mark_row(s["id"], written(cfg), 5), "component_id": new_uuid()}
    response = save_marks(admin, ex1["id"], cfg, [bad_comp])
    assert response.status_code == 409
    bad_student = {**mark_row(s["id"], written(cfg), 5), "student_id": new_uuid()}
    assert save_marks(admin, ex1["id"], cfg, [bad_student]).status_code == 409


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A23")
def test_unknown_config_not_found(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    body = {"exam_id": ex1["id"], "subject_config_id": new_uuid(), "marks": [mark_row(world.students[0]["id"], written(cfg), 5)]}
    response = admin.post(f"/exams/{ex1['id']}/marks", json=body)
    assert response.status_code == 404
    assert response.json()["detail"].startswith("ExamSubjectConfig with id ")


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A24")
def test_save_requires_exam_id(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    body = {"subject_config_id": cfg["id"], "marks": [mark_row(world.students[0]["id"], written(cfg), 5)]}
    assert admin.post(f"/exams/{ex1['id']}/marks", json=body).status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A25")
def test_save_on_published_exam_allowed(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    admin.post(f"/exams/{ex1['id']}/publish")
    assert save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 5)]).status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A26")
def test_second_attempt_creates_new_rows(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s = world.students[0]
    first = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 30)]).json()
    second = save_marks(admin, ex1["id"], cfg, [mark_row(s["id"], written(cfg), 55)], attempt=2)
    assert second.status_code == 200, second.text
    rows = second.json()
    assert len(rows) == 2
    attempts = {r["attempt_number"]: r for r in rows}
    assert attempts[1]["id"] == first[0]["id"] and attempts[1]["marks_obtained"] == "30.00"
    assert attempts[2]["marks_obtained"] == "55.00"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A27")
def test_remarks_component_stores_remark_grade(admin, world, cleanup):
    rs = create_remark_set(admin, cleanup)
    payload = exam_payload(world, subjects=("eng",))
    payload["subject_configs"][0]["components"].append(
        {"component_name": "Grade", "entry_type": "remarks", "remark_grade_set_id": rs["id"], "sort_order": 2}
    )
    response = admin.post("/exams", json=payload)
    assert response.status_code == 201, response.text
    e = response.json()["exam_id"]
    cleanup.add(remove_exam, admin, e)
    cfg = get_configs(admin, e)[0]
    grade = comp_by_name(cfg, "Grade")
    saved = save_marks(admin, e, cfg, [mark_row(world.students[0]["id"], grade, None, remark="A")])
    assert saved.status_code == 200, saved.text
    assert saved.json()[0]["remark_grade"] == "A"
    cell = student_row(grid(admin, e, world, cfg).json(), world.students[0])["marks"][grade["id"]]
    assert cell["remark_grade"] == "A"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A28")
@pytest.mark.parametrize("role", ["staff", "student", "parent"])
def test_save_denied(role_clients, world, ex1, role):
    cfg = math_cfg(world, ex1)
    assert save_marks(role_clients[role], ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 5)]).status_code == 403


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A29")
@pytest.mark.parametrize("role", ["admin", "teacher"])
def test_save_allowed(role_clients, world, ex1, role):
    cfg = math_cfg(world, ex1)
    assert save_marks(role_clients[role], ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 5)]).status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A30")
def test_saving_marks_writes_no_audit(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    save_marks(admin, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 5)])
    assert audit_actions(admin, ex1["id"]) == []


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A31")
def test_template_download(teacher, world, ex1):
    cfg = math_cfg(world, ex1)
    response = get_template(teacher, ex1["id"], world, cfg)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(XLSX_TYPE)
    rows = read_xlsx(response.content)
    assert rows[0][:3] == ["student_id", "Roll No", "Student Name"]
    assert "Written (Max: 80.00)" in rows[0] and "Oral (Max: 20.00)" in rows[0]
    assert rows[0][-1] == "Remarks"
    assert {r[0] for r in rows[1:]} == {s["id"] for s in world.students}


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A32")
def test_template_prefilled_with_marks_and_abs(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s0, s1 = world.students[0], world.students[1]
    save_marks(admin, ex1["id"], cfg, [mark_row(s0["id"], written(cfg), 61), mark_row(s1["id"], written(cfg), None, absent=True)])
    rows = read_xlsx(get_template(admin, ex1["id"], world, cfg).content)
    header = rows[0]
    col = header.index("Written (Max: 80.00)")
    by_student = {r[0]: r for r in rows[1:]}
    assert float(by_student[s0["id"]][col]) == 61
    assert by_student[s1["id"]][col] == "ABS"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A33")
def test_template_param_and_config_errors(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    params = tmpl_params(world, cfg)
    del params["subject_config_id"]
    assert admin.get(f"/exams/{ex1['id']}/marks/template", params=params).status_code == 422
    unknown = admin.get(
        f"/exams/{ex1['id']}/marks/template",
        params={"class_id": world.class_id, "section_id": world.sections["a"], "subject_config_id": new_uuid()},
    )
    assert unknown.status_code == 404
    assert unknown.json()["detail"] == "Subject config not found"


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A34")
@pytest.mark.parametrize("role", ["admin", "staff"])
def test_template_readers(role_clients, world, ex1, role):
    assert get_template(role_clients[role], ex1["id"], world, math_cfg(world, ex1)).status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A35")
@pytest.mark.parametrize("role", ["student", "parent"])
def test_template_denied(role_clients, world, ex1, role):
    assert get_template(role_clients[role], ex1["id"], world, math_cfg(world, ex1)).status_code == 403


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A36")
def test_upload_filled_template(teacher, admin, world, ex1):
    cfg = math_cfg(world, ex1)
    template = get_template(teacher, ex1["id"], world, cfg).content
    values = {s["id"]: [60 + i, 10 + i] for i, s in enumerate(world.students)}
    response = upload(teacher, ex1["id"], world, cfg, fill_template(template, values))
    assert response.status_code == 200, response.text
    assert response.json() == {"status": "done", "written": 6, "errors": [], "total_rows": 3}
    rows = grid(admin, ex1["id"], world, cfg).json()
    for i, s in enumerate(world.students):
        marks = student_row(rows, s)["marks"]
        assert marks[written(cfg)["id"]]["marks_obtained"] == 60 + i
        assert marks[oral(cfg)["id"]]["marks_obtained"] == 10 + i


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A37")
def test_upload_skips_abs_empty_and_text(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    template = get_template(admin, ex1["id"], world, cfg).content
    s0, s1, s2 = world.students
    values = {s0["id"]: [55, "ABS"], s1["id"]: [None, "abc"], s2["id"]: [70, None]}
    response = upload(admin, ex1["id"], world, cfg, fill_template(template, values))
    assert response.status_code == 200, response.text
    assert response.json()["written"] == 2
    assert response.json()["errors"] == []
    rows = grid(admin, ex1["id"], world, cfg).json()
    assert student_row(rows, s0)["marks"][written(cfg)["id"]]["marks_obtained"] == 55
    assert oral(cfg)["id"] not in student_row(rows, s0)["marks"] or student_row(rows, s0)["marks"][oral(cfg)["id"]]["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A38")
def test_upload_above_max_rejected(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    template = get_template(admin, ex1["id"], world, cfg).content
    s0, s1, _ = world.students
    values = {s0["id"]: [50, None], s1["id"]: [81, None]}
    response = upload(admin, ex1["id"], world, cfg, fill_template(template, values))
    assert response.status_code == 422
    rows = grid(admin, ex1["id"], world, cfg).json()
    cell = student_row(rows, s0)["marks"].get(written(cfg)["id"])
    assert cell is None or cell["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A39")
def test_upload_number_clears_absence(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s0 = world.students[0]
    save_marks(admin, ex1["id"], cfg, [mark_row(s0["id"], written(cfg), None, absent=True)])
    template = get_template(admin, ex1["id"], world, cfg).content
    response = upload(admin, ex1["id"], world, cfg, fill_template(template, {s0["id"]: [44, None]}))
    assert response.status_code == 200
    cell = student_row(grid(admin, ex1["id"], world, cfg).json(), s0)["marks"][written(cfg)["id"]]
    assert cell["is_absent"] is False and cell["marks_obtained"] == 44


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A40")
def test_upload_row_errors(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    s0 = world.students[0]
    content = build_xlsx(
        ["student_id", "Roll No", "Student Name", "Written (Max: 80.00)", "Oral (Max: 20.00)", "Remarks"],
        [[None, "1", "NoId", 50, 10, None], ["not-a-uuid", "2", "BadId", 50, 10, None], [s0["id"], "3", "Ok", 66, None, None]],
    )
    response = upload(admin, ex1["id"], world, cfg, content)
    assert response.status_code == 200, response.text
    body = response.json()
    assert any(e.startswith("Row 2: missing student_id") for e in body["errors"])
    assert any(e.startswith("Student not-a-uuid") for e in body["errors"])
    assert body["written"] == 1
    cell = student_row(grid(admin, ex1["id"], world, cfg).json(), s0)["marks"][written(cfg)["id"]]
    assert cell["marks_obtained"] == 66


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A41")
def test_upload_renamed_headers_ignored(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    content = build_xlsx(["student_id", "Roll No", "Student Name", "Written", "Oral"], [[world.students[0]["id"], "1", "A", 50, 10]])
    response = upload(admin, ex1["id"], world, cfg, content)
    assert response.status_code == 200
    assert response.json()["written"] == 0


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A42")
def test_upload_rejects_non_excel(admin, world, ex1):
    cfg = math_cfg(world, ex1)
    csv = upload(admin, ex1["id"], world, cfg, b"student_id,Written\n1,2\n", name="m.csv", ctype="text/csv")
    assert csv.status_code == 400
    assert csv.json()["detail"].startswith("Invalid Excel file")
    junk = upload(admin, ex1["id"], world, cfg, b"\x00\x01\x02random", name="m.xlsx")
    assert junk.status_code == 400


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A43")
@pytest.mark.parametrize("missing", ["class_id", "section_id", "subject_config_id"])
def test_upload_missing_params(admin, world, ex1, missing):
    cfg = math_cfg(world, ex1)
    params = tmpl_params(world, cfg)
    del params[missing]
    response = upload(admin, ex1["id"], world, cfg, build_xlsx(["student_id"], []), params=params)
    assert response.status_code == 422


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A44")
@pytest.mark.parametrize("role", ["staff", "student", "parent"])
def test_upload_denied(role_clients, world, ex1, role):
    cfg = math_cfg(world, ex1)
    response = upload(role_clients[role], ex1["id"], world, cfg, build_xlsx(["student_id"], []))
    assert response.status_code == 403


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A45")
def test_upload_skips_delegation_check(admin, teacher, world, ex1, logins):
    from api_tests.exam.test_f10_permissions import grant

    grant(admin, ex1["id"], logins["admin"]["user"]["id"])
    cfg = math_cfg(world, ex1)
    response = upload(teacher, ex1["id"], world, cfg, build_xlsx(["student_id"], []))
    assert response.status_code == 200


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A46")
def test_marks_unauthenticated(anon, world, ex1):
    cfg = math_cfg(world, ex1)
    e = ex1["id"]
    assert grid(anon, e, world, cfg).status_code == 401
    assert save_marks(anon, e, cfg, [mark_row(world.students[0]["id"], written(cfg), 5)]).status_code == 401
    assert get_template(anon, e, world, cfg).status_code == 401
    assert upload(anon, e, world, cfg, build_xlsx(["student_id"], [])).status_code == 401


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A47")
def test_marks_tenant_isolation(admin, tenant_b, world, ex1):
    cfg = math_cfg(world, ex1)
    e = ex1["id"]
    seen = grid(tenant_b, e, world, cfg)
    assert seen.status_code == 200
    assert not {r["student_id"] for r in seen.json()} & {s["id"] for s in world.students}
    response = save_marks(tenant_b, e, cfg, [mark_row(world.students[0]["id"], written(cfg), 5)])
    assert response.status_code in (404, 409)
    rows = grid(admin, e, world, cfg).json()
    cell = student_row(rows, world.students[0])["marks"].get(written(cfg)["id"])
    assert cell is None or cell["marks_obtained"] is None


@pytest.mark.api
@pytest.mark.tc("TC-EXM-11-A48")
def test_marks_cschema_mismatch(foreign, world, ex1):
    cfg = math_cfg(world, ex1)
    response = save_marks(foreign, ex1["id"], cfg, [mark_row(world.students[0]["id"], written(cfg), 5)])
    assert response.status_code == 403
