import datetime
import io

import openpyxl
import pytest

from api_tests.staff.helpers import (
    TEMP_PASSWORD,
    bulk_row,
    email_address,
    expected_status,
    granted,
    make_workbook,
    phone_number,
    remove_staff_by_id,
    upload_workbook,
)
from api_tests.support import QA_TENANT, Api, unique

UPLOAD = "/staff/enrollment/bulk-upload"
TEMPLATE = "/staff/enrollment/bulk-upload/template"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def upload(admin, cleanup, rows, headers=None, sheet="Staff Admission"):
    response = upload_workbook(admin, make_workbook(rows, headers, sheet))
    if response.status_code == 200:
        for created in response.json()["created"]:
            cleanup.add(remove_staff_by_id, admin, created["staff_id"])
    return response


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A01")
def test_template_download(admin):
    response = admin.get(TEMPLATE)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX)
    assert "staff_bulk_upload_template.xlsx" in response.headers["content-disposition"]
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Staff Admission" in workbook.sheetnames
    header = [cell.value for cell in workbook["Staff Admission"][1] if cell.value]
    assert len(header) == 33
    assert {"First Name", "Phone", "Address", "Designation", "Role"} <= set(header)


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A02")
def test_template_lists_current_designations(admin, make_designation):
    designation = make_designation(unique("stf_cook_"))
    workbook = openpyxl.load_workbook(io.BytesIO(admin.get(TEMPLATE).content))
    values = {cell.value for row in workbook["Lists"].iter_rows() for cell in row if cell.value}
    assert designation["title"] in values


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A03")
def test_upload_three_valid_rows(admin, cleanup, make_designation, logins):
    designation = make_designation()
    first, second, third = unique("stf_b_"), unique("stf_b_"), unique("stf_b_")
    rows = [
        bulk_row(first, level="Graduation", degree="B.Sc"),
        bulk_row(second, role="Teacher"),
        bulk_row(third, designation=designation["title"]),
    ]
    response = upload(admin, cleanup, rows)
    assert response.status_code == 200
    body = response.json()
    assert len(body["created"]) == 3 and body["errors"] == [] and body["total_rows"] == 3
    by_name = {c["name"]: c for c in body["created"]}
    assert [c["row"] for c in body["created"]] == [2, 3, 4]
    one = admin.get(f"/staff/enrollment/{by_name[first]['staff_id']}").json()
    assert [q["name"] for q in one["qualifications"]] == ["B.Sc"]
    two = admin.get(f"/staff/enrollment/{by_name[second]['staff_id']}").json()
    assert admin.get(f"/admin/users/{two['user_id']}").json()["role_id"] == logins["teacher"]["role"]["id"]
    three = admin.get(f"/staff/enrollment/{by_name[third]['staff_id']}").json()
    assert three["designation_id"] == designation["id"]


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A04")
def test_mixed_rows_report_errors(admin, cleanup):
    existing = email_address()
    seed = upload(admin, cleanup, [bulk_row(unique("stf_b_"), email=existing)])
    assert len(seed.json()["created"]) == 1
    rows = [
        bulk_row(unique("stf_b_")),
        bulk_row(unique("stf_b_"), email=existing),
        bulk_row(unique("stf_b_"), designation="stf_no_such_designation"),
        bulk_row(unique("stf_b_"), phone=""),
    ]
    response = upload(admin, cleanup, rows)
    assert response.status_code == 200
    body = response.json()
    assert len(body["created"]) == 1 and body["total_rows"] == 4
    assert body["errors"][0] == f"Row 3: A user with email '{existing}' already exists"
    assert body["errors"][1] == "Row 4: Designation 'stf_no_such_designation' not found"
    assert body["errors"][2] == "Row 5: Phone is required"


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A05")
def test_unknown_role_row(admin, cleanup):
    response = upload(admin, cleanup, [bulk_row(unique("stf_b_"), role="Nonexistent")])
    assert response.json()["created"] == []
    assert response.json()["errors"] == ["Row 2: Role 'Nonexistent' not found"]


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A06")
def test_qualification_level_without_degree(admin, cleanup):
    response = upload(admin, cleanup, [bulk_row(unique("stf_b_"), level="Graduation")])
    assert response.json()["errors"] == ["Row 2: Qualification Level and Degree/Course must be provided together"]


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A07")
def test_invalid_qualification_level(admin, cleanup):
    response = upload(admin, cleanup, [bulk_row(unique("stf_b_"), level="Masters", degree="X")])
    error = response.json()["errors"][0]
    assert error.startswith("Row 2: Invalid Qualification Level 'Masters'. Must be one of:")
    assert response.json()["created"] == []


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A08")
def test_bad_email_and_bad_gender(admin, cleanup):
    rows = [bulk_row(unique("stf_b_"), email="bad"), bulk_row(unique("stf_b_"), gender="X")]
    errors = upload(admin, cleanup, rows).json()["errors"]
    assert errors[0].startswith("Row 2: email - ")
    assert errors[1].startswith("Row 3: Invalid gender value: X.")


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A09")
def test_duplicate_email_inside_one_file(admin, cleanup):
    email = email_address()
    response = upload(admin, cleanup, [bulk_row(unique("stf_b_"), email=email), bulk_row(unique("stf_b_"), email=email)])
    body = response.json()
    assert len(body["created"]) == 1
    assert body["errors"] == [f"Row 3: A user with email '{email}' already exists"]


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A10")
def test_duplicate_phone_inside_one_file(admin, cleanup):
    phone = phone_number()
    response = upload(admin, cleanup, [bulk_row(unique("stf_b_"), phone=phone), bulk_row(unique("stf_b_"), phone=phone)])
    body = response.json()
    assert len(body["created"]) == 1
    assert len(body["errors"]) == 1
    assert body["errors"][0].startswith("Row 3: Error creating staff enrollment")


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A11")
@pytest.mark.parametrize("filename", ["staff.csv", "staff.txt"])
def test_wrong_extension(admin, filename):
    response = admin.post(UPLOAD, files={"file": (filename, b"a,b", "text/plain")})
    assert response.status_code == 400
    assert response.json()["detail"] == "File must be an Excel (.xlsx/.xls) file"


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A12")
def test_garbage_xlsx(admin):
    response = admin.post(UPLOAD, files={"file": ("staff.xlsx", b"garbage bytes", XLSX)})
    assert response.status_code == 400
    assert response.json()["detail"].startswith("Invalid Excel file")


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A13")
def test_missing_address_header(admin):
    data = make_workbook([], headers=["First Name", "Phone"])
    response = upload_workbook(admin, data)
    assert response.status_code == 400
    assert response.json()["detail"] == "Missing required column(s): Address"


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A14")
def test_header_only_sheet(admin):
    response = upload_workbook(admin, make_workbook([]))
    assert response.status_code == 200
    assert response.json() == {"created": [], "errors": [], "total_rows": 0}


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A15")
def test_date_formats(admin, cleanup):
    rows = [
        bulk_row(unique("stf_b_"), dob="25-12-1990"),
        bulk_row(unique("stf_b_"), dob=datetime.datetime(1990, 12, 25)),
    ]
    body = upload(admin, cleanup, rows).json()
    assert len(body["created"]) == 2
    for created in body["created"]:
        assert admin.get(f"/staff/enrollment/{created['staff_id']}").json()["date_of_birth"] == "1990-12-25"


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A16")
def test_uploaded_staff_must_change_password(admin, cleanup, academic_year_id):
    phone = phone_number()
    body = upload(admin, cleanup, [bulk_row(unique("stf_b_"), phone=phone)]).json()
    assert len(body["created"]) == 1
    client = Api(tenant_header=QA_TENANT)
    try:
        response = client.post(
            "/auth/login", json={"username": phone, "password": TEMP_PASSWORD, "academic_year_id": academic_year_id}
        )
    finally:
        client.close()
    assert response.status_code == 200
    assert response.json()["requires_password_change"] is True


@pytest.mark.api
@pytest.mark.parametrize(
    "role",
    [
        pytest.param("staff", marks=pytest.mark.tc("TC-STF-08-A17")),
        pytest.param("teacher", marks=pytest.mark.tc("TC-STF-08-A17")),
        pytest.param("student", marks=pytest.mark.tc("TC-STF-08-A17")),
        pytest.param("parent", marks=pytest.mark.tc("TC-STF-08-A17")),
    ],
)
def test_bulk_upload_denied_for_non_admin(role_clients, logins, role):
    assert not granted(logins, role, "staff", "create")
    client = role_clients[role]
    name = unique("stf_den_")
    response = upload_workbook(client, make_workbook([bulk_row(name)]))
    assert response.status_code == 403
    assert client.get(TEMPLATE).status_code == expected_status(logins, role, "staff", "create")
    listed = [s for s in role_clients["admin"].get("/staff/enrollments").json() if s["first_name"] == name]
    assert listed == []


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A18")
def test_bulk_upload_requires_token(anon):
    assert anon.get(TEMPLATE).status_code == 401
    assert upload_workbook(anon, make_workbook([bulk_row("x")])).status_code == 401


@pytest.mark.api
@pytest.mark.tc("TC-STF-08-A19")
def test_bulk_upload_tenant_isolation(admin, tenant_b, cleanup):
    title = unique("stf_bonly_")
    created = tenant_b.post("/staff/designations/", json={"title": title})
    assert created.status_code == 201
    cleanup.delete_later(tenant_b, f"/staff/designations/{created.json()['id']}")
    name = unique("stf_b_")
    response = upload(admin, cleanup, [bulk_row(name, designation=title)])
    assert response.json()["created"] == []
    assert response.json()["errors"] == [f"Row 2: Designation '{title}' not found"]
    assert name not in [s["first_name"] for s in tenant_b.get("/staff/enrollments").json()]
